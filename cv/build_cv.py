# -*- coding: utf-8 -*-
"""Generate bilingual (EN + CN) CV for Jenny Wu targeting VIVAIA Social Media Marketing Manager/Supervisor."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x8B, 0x6F, 0x4E)   # warm taupe, VIVAIA-adjacent neutral
DARK = RGBColor(0x2B, 0x2B, 0x2B)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)

doc = Document()

# page margins
for s in doc.sections:
    s.top_margin = Cm(1.6)
    s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(1.9)
    s.right_margin = Cm(1.9)

# base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.font.color.rgb = DARK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.12


def set_east_asia(run, font="微软雅黑"):
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font)


def hr(paragraph, color="C9B99B", size="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def name_header(text_en, text_cn):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text_en)
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = DARK
    set_east_asia(r)
    r2 = p.add_run("  " + text_cn)
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT
    set_east_asia(r2)
    return p


def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = ACCENT
    r.font.bold = True
    set_east_asia(r)
    return p


def contact_line(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY
    set_east_asia(r)
    hr(p)
    return p


def section(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    set_east_asia(r)
    hr(p, size="4")
    return p


def para(text, size=10, bold=False, italic=False, color=DARK, space_after=2, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_east_asia(r)
    return p


def job_header(company, role, date_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(company)
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = DARK
    set_east_asia(r)
    r2 = p.add_run("  |  " + date_str)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = GRAY
    set_east_asia(r2)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r3 = p2.add_run(role)
    r3.font.size = Pt(10)
    r3.font.bold = True
    r3.font.color.rgb = ACCENT
    set_east_asia(r3)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Cm(0.55)
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = DARK
        set_east_asia(r0)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = DARK
    set_east_asia(r)
    return p


# ============================================================
# ENGLISH VERSION
# ============================================================
name_header("JENNY WU", "巫婧怡")
subtitle("Social Media & Brand Marketing Manager — Global DTC Brands")
contact_line("Shenzhen, China  ·  [Phone]  ·  [Email]  ·  [LinkedIn / Portfolio]  ·  Working languages: English & Mandarin")

section("PROFILE")
para("Global social media marketer with 3+ years building overseas brands 0→1 across DTC e-commerce, lifestyle and consumer hardware. Pairs a Visual Communication Design degree with hands-on performance literacy — equally fluent in brand aesthetics and CPL/CTR spreadsheets. Currently leading Instagram-first content strategy, Meta Ads experimentation and creator marketing for a North American kids' music-tech brand. Builds AI-augmented content pipelines that scale output 10x more cost-efficiently without diluting brand visual standards.")

section("CORE STRENGTHS")
bullet("content pillars, Reels systems, grid visual language, community growth", "Instagram-first global social strategy — ")
bullet("B.A. in Visual Communication Design; art direction, briefs and QA for in-house teams, creators and outsourced designers", "Brand visual direction — ")
bullet("prompt engineering with brand-safety constraints (Jimeng AI, image/video generation); asset costs cut to ~$0.1/image and ~$5/video", "AI-empowered content production — ")
bullet("sourcing, outreach, briefing and performance tracking across lifestyle / education / family verticals; relationship-first DM philosophy", "Creator & influencer marketing 0→1 — ")
bullet("Meta Ads Manager, Pixel/CAPI implementation, clean A/B test design, CPL/CTR-driven iteration", "Performance marketing literacy — ")
bullet("native-level English copywriting for NA-facing deliverables; cross-functional and agency management", "Bilingual campaign delivery (EN/CN) — ")

section("PROFESSIONAL EXPERIENCE")

job_header("MuseKey (playmusekey.com) — Screen-free AI Music Creation Device for Kids, North America",
           "Content Operations Lead — Social Media & Creator Marketing", "[MM/YYYY] – Present")
para("Own global social (Instagram-first), paid social experimentation and influencer marketing for a kids' consumer hardware brand entering the US market; report directly to the founder.", italic=True, color=GRAY, size=9.5)
bullet("designed the PLAY · CREATE · GROW content pillar framework and the brand's visual language across Reels and static formats; manage weekly asset calendars and cross-team asset ledgers", "Built the Instagram content system from zero — ")
bullet("first baseline flight delivered 68 email leads at ~¥15.8 ($2.2) blended CPL; top creative hit 9.85% CTR. Surfaced a counter-intuitive audience insight (women 55+ over-indexing on FB) and designed a clean IG-vs-FB platform A/B test to validate it", "Ran Meta Ads as structured experiments — ")
bullet("implemented Meta Pixel + Conversions API, Lead event mapping, domain verification and privacy compliance on the DTC site", "Shipped full-funnel tracking — ")
bullet("built creator sourcing and outreach playbooks across parenting, children's education and creative-toy verticals; relationship-first DM framework prioritizing authentic connection over transactional pitching", "Influencer marketing 0→1 — ")
bullet("standardized prompt-engineering workflows (Jimeng AI) with anti-defect visual constraints; strategy kept in-house, execution delegated to outsourced editors via brief + acceptance-criteria model", "AI-augmented production pipeline — ")
bullet("coordinate a 4-person cross-functional team plus 3 outsourced squads (AI video, community listening, data tracking)", "Team leadership — ")

job_header("TopJob — North America DTC Ergonomic Chair Brand",
           "Social Media & Creator Marketing", "[MM/YYYY] – [MM/YYYY]")
para("Built the brand's North America creator program from scratch and ran Black Friday–Christmas growth campaigns end-to-end.", italic=True, color=GRAY, size=9.5)
bullet("sourced and evaluated 50 NA creators into a tiered pyramid (top/mid/nano KOCs); standardized the full SOP — outreach, briefs, content review, publishing and data tracking", "Creator system 0→1 — ")
bullet("17 published collaborations delivered 154K impressions and 3,218 engagements at 2.09% average ER — 4x+ the home-category Instagram benchmark (0.2–0.5%); zero-boost creators averaged 6.02% ER vs 1.02% for boosted posts", "Outperforming engagement — ")
bullet("orchestrated 17 creators across a 4-stage content layering model (seeding → explaining → converting → review); social ads ROI 4.14, brand exposure +12%, Q4 creator ROI doubled vs Q3", "Black Friday–Christmas campaign — ")
bullet("appeared on-camera producing TikTok ad creatives with multi-hook testing; fed winning angles (e.g. assembly convenience) back into creator briefs in real time", "Hands-on content production — ")
bullet("wrote scene-based briefs (\"film how your back feels after a full day\") instead of feature readouts, turning product specs into user decision answers", "Brief methodology — ")

job_header("DOME Group — Premium F&B & Lifestyle Brand Group, Shanghai",
           "Overseas Content & Growth Marketing", "[MM/YYYY] – [MM/YYYY]")
para("Led 0→1 overseas content ecosystem for a new premium dining brand targeting expats, business travelers and international tourists; simultaneously optimized social growth for 3 mature group brands.", italic=True, color=GRAY, size=9.5)
bullet("built the Instagram + website + content matrix from zero; ran user research via IG keyword search, YouTube travel communities and Facebook groups to define international guest personas", "0→1 overseas brand building — ")
bullet("identified that overseas guests couldn't book (WeChat-only reservation flow); escalated and shipped a direct web-booking entry, removing the key conversion breakpoint", "Fixed the conversion gap — ")
bullet("ran a 3-phase creator campaign (seeding → explaining → converting) with 20+ creators; ¥27,810 directly attributed revenue on ¥19,400 spend — direct ROI 1.43; 323 new followers; weekly foot traffic +20%", "New Year's Eve campaign — ")
bullet("chef-collaboration content series: 20 KOL pieces, ¥21,000 package sales, 3.58% booking-link CTR", "Tuna Cutting Ceremony event — ")
bullet("in 6 months: 700K+ cumulative brand exposure, +35% foot traffic, blended creator marketing ROI ~1.2:1 including content asset reuse", "Overall results — ")

section("EDUCATION")
para("B.A. in Visual Communication Design — Lingnan Normal University (岭南师范学院)", bold=False)

section("TOOLBOX")
para("Meta Ads Manager · Pixel / Conversions API · Instagram · Facebook · TikTok · Pinterest · Influencity · FastMoss · Jimeng AI (即梦) & image/video prompt engineering · CapCut brief workflows · Feishu · Google Workspace · ClickUp")

# ============================================================
# CHINESE VERSION
# ============================================================
doc.add_page_break()

name_header("巫婧怡", "JENNY WU")
subtitle("社交媒体与品牌营销经理 — 全球化 DTC 品牌")
contact_line("中国 · 深圳  ·  [电话]  ·  [邮箱]  ·  [LinkedIn / 作品集]  ·  工作语言：中文（母语）、英文（母语级文案产出）")

section("个人简介 PROFILE")
para("3 年以上海外品牌 0→1 社媒营销经验，覆盖 DTC 电商、生活方式与消费硬件品类。视觉传达设计专业出身，兼具品牌视觉审美与投放数据素养（CPL/CTR 驱动迭代）。现任北美儿童音乐科技品牌内容运营负责人，主导 Instagram 为核心的内容策略、Meta 广告实验与达人营销。搭建 AI 赋能的内容生产管线，以约 1/10 的成本实现品牌级视觉素材的规模化产出。")

section("核心优势 CORE STRENGTHS")
bullet("内容支柱框架、Reels 体系、网格视觉语言、社群增长", "Instagram 为核心的全球社媒策略 — ")
bullet("视觉传达设计学士；为内部团队、达人与外包设计提供美术指导、Brief 与验收", "品牌视觉把控 — ")
bullet("带品牌安全约束的提示词工程（即梦 AI、图像/视频生成）；素材成本降至约 ¥0.8/图、¥31–39/视频", "AI 赋能内容生产 — ")
bullet("覆盖生活方式/教育/家庭赛道的达人筛选、建联、Brief 与数据追踪；坚持关系优先的触达哲学", "达人营销 0→1 — ")
bullet("Meta Ads Manager、Pixel/CAPI 埋点、纯净 A/B 实验设计、CPL/CTR 驱动迭代", "效果营销素养 — ")
bullet("面向北美的母语级英文文案；跨职能团队与外包/代理商管理", "中英双语战役交付 — ")

section("工作经历 PROFESSIONAL EXPERIENCE")

job_header("MuseKey（playmusekey.com）— 面向北美的儿童无屏 AI 音乐创作硬件",
           "内容运营负责人 — 社媒与达人营销", "[起始月份] – 至今")
para("负责品牌进入美国市场的全球社媒（Instagram 为核心）、付费投放实验与达人营销，直接向创始人汇报。", italic=True, color=GRAY, size=9.5)
bullet("设计 PLAY · CREATE · GROW 内容支柱框架与品牌视觉语言，覆盖 Reels 与静态素材；管理周度素材排期与跨团队资产台账", "从零搭建 Instagram 内容体系 — ")
bullet("首轮基线投放以 ¥1,077 获得 68 个邮箱线索，混合 CPL ~¥15.8（$2.2）；最优素材 CTR 9.85%。发现反直觉人群洞察（FB 端 55+ 女性占比过半），并设计 IG vs FB 纯净平台 A/B 实验验证", "以实验方法运营 Meta 广告 — ")
bullet("完成 Meta Pixel + CAPI 接入、Lead 事件映射、域名验证与隐私合规上线", "打通全链路数据追踪 — ")
bullet("围绕 parenting、儿童教育、创意玩具三大赛道建立达人筛选与触达话术体系，坚持真诚连接优先、低交易感的 DM 框架", "达人营销 0→1 — ")
bullet("沉淀含反缺陷视觉约束的提示词工程流程（即梦 AI）；策略 in-house、执行外包（Brief + 验收标准模式）", "AI 赋能生产管线 — ")
bullet("统筹 4 人跨职能团队及 3 支外包小组（AI 视频、社群监听、数据打点）", "团队管理 — ")

job_header("TopJob — 北美 DTC 人体工学椅品牌",
           "海外社媒与达人营销", "[起止月份]")
para("从零搭建品牌北美达人合作体系，并端到端负责黑五–圣诞购物季增长战役。", italic=True, color=GRAY, size=9.5)
bullet("筛选评估 50 位北美达人，形成头部/腰部/尾部 KOC 金字塔结构；沉淀建联、Brief、内容审核、发布与数据追踪全流程 SOP", "达人体系 0→1 — ")
bullet("17 个已发布合作累计曝光 15.4 万、互动 3,218 次，平均互动率 2.09%，超家居行业 Instagram 基准（0.2%–0.5%）4 倍以上；零推流达人平均互动率 6.02%，远高于推流内容的 1.02%", "互动表现远超行业基准 — ")
bullet("按「种草→解释→转化→复盘」四阶段编排 17 位达人节奏；社媒广告 ROI 4.14，品牌曝光提升 12%，Q4 达人合作 ROI 较 Q3 翻倍", "黑五–圣诞战役 — ")
bullet("亲自出镜拍摄 TikTok 广告素材，多 Hook 快速迭代测试，将跑出的最优角度（如组装便利性）实时回写到达人 Brief", "亲自动手的内容生产 — ")
bullet("Brief 不写「介绍腰托功能」，而写「拍你坐了一整天的真实感受」，把产品参数翻译成用户决策答案", "场景化 Brief 方法论 — ")

job_header("DOME Group — 高端餐饮与生活方式品牌集团（上海）",
           "海外内容与增长营销", "[起止月份]")
para("主导集团新高端餐饮品牌的海外内容体系 0→1 搭建，目标客群为在沪外籍人士、商务人群与国际游客；同时负责 3 个成熟品牌的社媒增长优化。", italic=True, color=GRAY, size=9.5)
bullet("从零搭建 Instagram + 官网 + 内容矩阵；通过 IG 关键词搜索、YouTube 旅行社群与 Facebook 群组调研定义国际用户画像", "海外品牌 0→1 — ")
bullet("发现海外用户因仅支持微信预约而无法下单，持续推动上线官网直接预订入口，彻底打通转化断点", "修复转化断点 — ")
bullet("以「种草→解释→转化」三阶段打法统筹 20+ 位达人；达人直接带来营收 ¥27,810、总投入 ¥19,400，直接 ROI 1.43；新增粉丝 323，周到店客流 +20%", "跨年夜节点战役 — ")
bullet("主厨合作内容系列：20 位 KOL 内容产出，套餐销售 ¥21,000，预订链接 CTR 3.58%", "开鱼仪式活动 — ")
bullet("6 个月累计品牌曝光 70 万+，到店客流增长 35%，含内容资产复用的达人营销综合 ROI 约 1.2:1", "整体成果 — ")

section("教育背景 EDUCATION")
para("岭南师范学院 — 视觉传达设计 学士")

section("工具栈 TOOLBOX")
para("Meta Ads Manager · Pixel / CAPI · Instagram · Facebook · TikTok · Pinterest · Influencity · FastMoss · 即梦 AI 及图像/视频提示词工程 · CapCut Brief 流程 · 飞书 · Google Workspace · ClickUp")

out = "/Users/mima0000/WorkBuddy/2026-09-21-17-00-41/JennyWu_CV_VIVAIA_v2.docx"
doc.save(out)
print("saved:", out)
